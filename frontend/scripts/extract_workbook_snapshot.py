"""Create frontend preview data from a read-only copy of the IGP workbook.

Usage: python scripts/extract_workbook_snapshot.py /path/to/workbook.xlsx
This script never writes to the source workbook.
"""
import json
import sys
from pathlib import Path

import openpyxl

source = Path(sys.argv[1])
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
core = book["Core Data"]
funds = book["Partners Funds"]
credits = book["Credit Data"]
cash = book["Petty Cash Index"]
index = book["Partners Index"]

names = ["DSP", "SISU", "RAISE", "Anbu Illam", "Shalom", "Duncan", "HOL", "BHS", "JJS", "ISD"]
fund_rows = {"DSP": 5, "SISU": 6, "RAISE": 7, "Anbu Illam": 8, "Shalom": 9, "Duncan": 10, "HOL": 11, "JJS": 12, "BHS": 13, "ISD": 14}
count_columns = {"DSP": 3, "SISU": 4, "RAISE": 5, "Anbu Illam": 6, "Duncan": 7, "Shalom": 8, "HOL": 9, "JJS": 10, "ISD": 11, "BHS": 12}


def number(value):
    return round(float(value), 6) if isinstance(value, (int, float)) else 0


partners = []
for i, name in enumerate(names):
    monthly_row = fund_rows[name]
    count_col = count_columns[name]
    partners.append({
        "name": name,
        "budget": number(core.cell(122 + i, 7).value),
        "finalGrant": number(core.cell(122 + i, 8).value),
        "schedule": [number(funds.cell(monthly_row, c).value) for c in range(4, 16)],
        "received": [number(credits.cell(monthly_row, c).value) for c in range(4, 16)],
        "cs": number(index.cell(4, count_col).value),
        "ecc": number(index.cell(5, count_col).value),
        "ea": number(index.cell(8, count_col).value),
        "socialWorkers": number(index.cell(9, count_col).value),
        "psw": number(index.cell(10, count_col).value),
        "apeTarget": number(index.cell(12, count_col).value),
    })

expenses = [
    {"category": str(cash.cell(r, 3).value).strip(), "monthly": [number(cash.cell(r, c).value) for c in range(4, 16)]}
    for r in range(5, 24)
    if cash.cell(r, 3).value
]

snapshot = {
    "year": 2026,
    "source": source.name,
    "snapshotDate": "2026-09-26",
    "months": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"],
    "partners": partners,
    "expenses": expenses,
    "workbookTotals": {"budget": number(core["G132"].value), "finalGrant": number(core["H132"].value)},
}

destination = Path(__file__).resolve().parents[1] / "src" / "data" / "workbookSnapshot.json"
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps(snapshot, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(f"Extracted {len(partners)} partners and {len(expenses)} expense categories to {destination}")
