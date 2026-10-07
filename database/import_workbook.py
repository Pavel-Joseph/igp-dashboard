"""Generate SQL for the four populated admin sheets. Read-only input workbook."""
from __future__ import annotations
import argparse
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import openpyxl

CODES = ("DSP", "SUSU", "RAISE", "AI", "DUNCAN", "SHALOM", "HOL",
         "JJS", "ISD", "BHS", "CHF", "KMVS")
NAMES = ("DSP", "SUSU", "RAISE", "Anbu Illam", "DUNCAN", "SHALOM", "HOL",
         "JJS", "ISD", "BHS", "CHF", "KMVS")
SCHEDULE_ROWS = {"DSP": 6, "SUSU": 7, "RAISE": 8, "AI": 9, "DUNCAN": 11,
                 "SHALOM": 10, "HOL": 12, "JJS": 13, "ISD": 15,
                 "BHS": 14, "CHF": 16, "KMVS": 17}
EXPECTED_SHEETS = {"Partners Grant'26", "Partners Index'26",
                   "Partners Schedule'26", "Petty Cash'26"}

def q(value: object) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"

def n(value: object, places: int = 2) -> str:
    if value is None:
        return "null"
    if not isinstance(value, (int, float, Decimal)):
        raise ValueError(f"Expected number, got {value!r}")
    scale = Decimal("1").scaleb(-places)
    return str(Decimal(str(value)).quantize(scale, rounding=ROUND_HALF_UP))

def partner_id(code: str) -> str:
    return f"(select id from public.partners where code={q(code)})"

def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    formulas = openpyxl.load_workbook(args.workbook, data_only=False, read_only=False)
    values = openpyxl.load_workbook(args.workbook, data_only=True, read_only=False)
    missing = EXPECTED_SHEETS - set(formulas.sheetnames)
    if missing:
        raise ValueError(f"Missing sheets: {sorted(missing)}")
    grant = values["Partners Grant'26"]
    index = values["Partners Index'26"]
    schedule = formulas["Partners Schedule'26"]
    schedule_values = values["Partners Schedule'26"]
    cash = values["Petty Cash'26"]
    lines = ["-- Generated from Finance Reports for 2026.xlsx; run after 001_admin_workbook.sql.",
             "-- Formula totals are recalculated via database queries, not imported.",
             "-- External workbook formulas have no verified amount; their cached value is retained separately.",
             "begin;"]
    counts = {"partners": 0, "grants": 0, "grant_items": 0,
              "metrics": 0, "schedules": 0, "petty_cash": 0}
    for code, name in zip(CODES, NAMES):
        lines.append(f"insert into public.partners(code,name) values({q(code)},{q(name)}) "
                     f"on conflict(code) do update set name=excluded.name;")
        counts["partners"] += 1
    for col, code in enumerate(CODES, start=3):
        lines.append(
            "insert into public.partner_grants(partner_id,fiscal_year,entered_total,"
            "prior_year_unutilised,entered_effective_grant) values("
            f"{partner_id(code)},2026,{n(grant.cell(14,col).value)},"
            f"{n(grant.cell(15,col).value)},{n(grant.cell(16,col).value)}) "
            "on conflict(partner_id,fiscal_year) do update set "
            "entered_total=excluded.entered_total,"
            "prior_year_unutilised=excluded.prior_year_unutilised,"
            "entered_effective_grant=excluded.entered_effective_grant;")
        counts["grants"] += 1
        for row in range(4,14):
            value = grant.cell(row,col).value
            if value is None:
                continue
            category = grant.cell(row,2).value
            lines.append(
                "insert into public.partner_grant_items(partner_id,fiscal_year,category,amount) "
                f"values({partner_id(code)},2026,{q(category)},{n(value)}) "
                "on conflict(partner_id,fiscal_year,category) do update set amount=excluded.amount;")
            counts["grant_items"] += 1
        for row in range(4,14):
            value = index.cell(row,col).value
            if value is None:
                continue
            metric = index.cell(row,2).value
            unit = "target_people" if metric == "APE" else ("unspecified" if metric == "Admin" else "people")
            lines.append(
                "insert into public.partner_metrics(partner_id,fiscal_year,metric,unit,value) "
                f"values({partner_id(code)},2026,{q(metric)},{q(unit)},{n(value,3)}) "
                "on conflict(partner_id,fiscal_year,metric) do update set "
                "unit=excluded.unit,value=excluded.value;")
            counts["metrics"] += 1
        row = SCHEDULE_ROWS[code]
        for month, col_month in enumerate(range(4,16),start=1):
            cell = schedule.cell(row,col_month)
            if cell.value is None:
                continue
            if cell.data_type == "f":
                amount = "null"
                formula = q(cell.value)
                cached = n(schedule_values.cell(row,col_month).value)
                status = "unverified_external_formula"
            else:
                amount = n(cell.value)
                formula = cached = "null"
                status = "verified"
            lines.append(
                "insert into public.partner_schedules(partner_id,fiscal_year,month,amount,"
                "source_formula,cached_amount,verification_status) "
                f"values({partner_id(code)},2026,{month},{amount},{formula},{cached},{q(status)}) "
                "on conflict(partner_id,fiscal_year,month) do update set "
                "amount=excluded.amount,source_formula=excluded.source_formula,"
                "cached_amount=excluded.cached_amount,"
                "verification_status=excluded.verification_status;")
            counts["schedules"] += 1
    for row in range(5,25):
        category = cash.cell(row,3).value
        if category is None:
            continue
        for month, col in enumerate(range(4,16),start=1):
            value = cash.cell(row,col).value
            if value is None:
                continue
            lines.append(
                "insert into public.petty_cash_monthly(fiscal_year,month,category,amount) "
                f"values(2026,{month},{q(category)},{n(value)}) "
                "on conflict(fiscal_year,month,category) do update set amount=excluded.amount;")
            counts["petty_cash"] += 1
    lines.append("commit;")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("\n".join(lines)+"\n",encoding="utf-8")
    print("Generated:", args.output)
    print("Rows:", counts)
    print("Unverified schedule formulas:",
          [(schedule.cell(r,c).coordinate,schedule.cell(r,c).value)
           for r in range(6,18) for c in range(4,16)
           if schedule.cell(r,c).data_type == "f"])

if __name__ == "__main__":
    main()
